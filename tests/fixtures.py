"""Builds one small document of each format, for the tests.

The files are generated (not committed), so the tests show exactly what each
document contains.
"""

from pathlib import Path


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = "Travel policy"
    document.add_heading("Travel policy", level=1)
    document.add_paragraph("Employees may book economy flights for trips under six hours.")
    document.add_heading("Hotels", level=2)
    document.add_paragraph("Hotels must cost less than 150 euros per night.")
    document.add_paragraph("Book through the travel portal", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "City"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Paris"
    table.cell(1, 1).text = "200"
    document.save(str(path))


def make_pptx(path: Path) -> None:
    import pptx

    presentation = pptx.Presentation()
    for title, body, notes in [
        ("Quarterly results", "Revenue grew 12 percent", "Mention the new customers."),
        ("Roadmap", "Launch the mobile app in spring", "Dates are not final."),
    ]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    presentation.save(str(path))


def make_xlsx(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["Product", "Price"])
    sheet.append(["Basic plan", 10])
    sheet.append(["Pro plan", 25.5])
    workbook.save(str(path))


def make_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for page_text in [
        "Warranty terms. The warranty covers manufacturing defects for two years.",
        "Returns. Products can be returned within thirty days of delivery.",
    ]:
        pdf.add_page()
        pdf.multi_cell(0, 10, page_text)
    pdf.output(str(path))


def make_all(folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    make_docx(folder / "travel.docx")
    make_pptx(folder / "results.pptx")
    make_xlsx(folder / "prices.xlsx")
    make_pdf(folder / "warranty.pdf")
    (folder / "notes.md").write_text("# Notes\n\n## Meeting\n\nWe agreed to ship on Friday.\n")
    (folder / "page.html").write_text(
        "<html><head><title>FAQ</title></head><body><nav>menu</nav><main>"
        "<h1>FAQ</h1><h2>Shipping</h2><p>Orders ship within two days.</p></main></body></html>"
    )
    return folder
