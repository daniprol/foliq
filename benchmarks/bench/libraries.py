"""One `extract(path) -> text` function per library, with the formats it reads.

Each library is used the way its README shows, with default options, asking for
Markdown when it can produce it. Competitors that are not Python packages
(pdftotext, pandoc, Apache Tika) are run through `tools.py`.
foliq is not here: it indexes whole folders, so `parsing.py` runs it
separately (see `run_foliq`).
"""

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from bench import tools


@dataclass
class Library:
    """A competitor: its name, the formats it reads, and how to extract text.

    `language` is what does the work (shown in the charts): PyMuPDF is a
    Python package, but its engine, MuPDF, is C. `requires` is the
    command a competitor needs when it is not a Python package; a competitor
    whose command is missing is skipped.
    """

    name: str
    formats: list[str]
    extract: Callable[[Path], str]
    language: str = "Python"
    requires: str | None = None

    def is_installed(self) -> bool:
        return self.requires is None or tools.is_installed(self.requires)

    @property
    def label(self) -> str:
        """The name shown in results and charts, e.g. `mutool (C)`."""
        return f"{self.name} ({self.language})"

    @property
    def folder(self) -> str:
        """A file name for this library's outputs (`ledongthuc/pdf` has a slash)."""
        return self.name.replace("/", "-")


def pymupdf(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as document:
        return "\n".join(page.get_text() for page in document)


def pypdfium2(path: Path) -> str:
    import pypdfium2

    document = pypdfium2.PdfDocument(path)
    try:
        return "\n".join(page.get_textpage().get_text_range() for page in document)
    finally:
        document.close()


def pypdf(path: Path) -> str:
    import pypdf

    reader = pypdf.PdfReader(path)
    return "\n".join(page.extract_text() for page in reader.pages)


def pdfplumber(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(path) as document:
        return "\n".join(page.extract_text() or "" for page in document.pages)


def markitdown(path: Path) -> str:
    from markitdown import MarkItDown

    return MarkItDown().convert(str(path)).text_content


def kreuzberg(path: Path) -> str:
    from kreuzberg import ExtractionConfig, extract_file_sync

    # No OCR, so every library reads only the text layer; Markdown output.
    config = ExtractionConfig(disable_ocr=True, output_format="markdown", use_cache=False)
    return extract_file_sync(path, config=config).content


def python_docx(path: Path) -> str:
    import docx

    document = docx.Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def python_pptx(path: Path) -> str:
    import pptx

    lines = []
    for slide in pptx.Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:
            lines.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(lines)


def openpyxl(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if any(value is not None for value in row):
                lines.append(" | ".join("" if value is None else str(value) for value in row))
    workbook.close()
    return "\n".join(lines)


def python_calamine(path: Path) -> str:
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_path(str(path))
    lines = []
    for name in workbook.sheet_names:
        lines.append(f"# {name}")
        for row in workbook.get_sheet_by_name(name).to_python(skip_empty_area=True):
            lines.append(" | ".join(str(value) for value in row))
    return "\n".join(lines)


def html2text(path: Path) -> str:
    import html2text

    return html2text.html2text(path.read_text(errors="replace"))


def trafilatura(path: Path) -> str:
    import trafilatura

    return trafilatura.extract(path.read_text(errors="replace"), output_format="markdown") or ""


def beautifulsoup(path: Path) -> str:
    from bs4 import BeautifulSoup

    return BeautifulSoup(path.read_text(errors="replace"), "html.parser").get_text("\n")


def pdftotext(path: Path) -> str:
    # Poppler, the PDF engine of most Linux desktops.
    return tools.run(["pdftotext", "-q", "-enc", "UTF-8", str(path), "-"])


def pandoc(path: Path) -> str:
    # Plain Markdown: without `-raw_html`, pandoc keeps divs, spans and tables as HTML.
    return tools.run(["pandoc", "--to", "gfm-raw_html", "--wrap", "none", str(path)])


# Per format: the libraries everyone knows, and the fastest ones (for a fair
# speed comparison). Others were tried and left out to keep the charts readable:
# pymupdf4llm, pdfminer.six (pdfplumber is built on it), mutool (same engine as
# PyMuPDF), pdf.js, pdf-extract, lopdf, mammoth, SheetJS, excelize,
# officeparser and extractous.
LIBRARIES = [
    # PDF
    Library("pymupdf", ["pdf"], pymupdf, language="C"),
    Library("pypdfium2", ["pdf"], pypdfium2, language="C++"),
    Library("pypdf", ["pdf"], pypdf),
    Library("pdfplumber", ["pdf"], pdfplumber),
    Library("pdftotext", ["pdf"], pdftotext, language="C++", requires="pdftotext"),
    # Office
    Library("python-docx", ["docx"], python_docx),
    Library("python-pptx", ["pptx"], python_pptx),
    Library("openpyxl", ["xlsx"], openpyxl),
    Library("python-calamine", ["xlsx"], python_calamine, language="Rust"),
    # HTML
    Library("beautifulsoup4", ["html"], beautifulsoup),
    Library("html2text", ["html"], html2text),
    Library("trafilatura", ["html"], trafilatura),
    Library("pandoc", ["docx", "html"], pandoc, language="Haskell", requires="pandoc"),
    # Many formats
    Library("markitdown", ["pdf", "docx", "pptx", "xlsx", "html"], markitdown),
    Library("kreuzberg", ["pdf", "docx", "pptx", "xlsx", "html"], kreuzberg, language="Rust"),
    Library("Apache Tika", ["pdf", "docx", "pptx", "xlsx"], tools.tika, language="Java", requires="tika"),
]

BY_NAME = {library.name: library for library in LIBRARIES}


def libraries_for(file_format: str) -> list[Library]:
    """The installed competitors that read this format; the missing ones are named."""
    found = []
    for library in LIBRARIES:
        if file_format not in library.formats:
            continue
        if library.is_installed():
            found.append(library)
        else:
            print(f"  skipping {library.name}: `{library.requires}` not found (see `python -m bench tools`)")
    return found
